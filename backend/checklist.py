"""Monthly checklist discovery, normalization, matching, and coverage tree data."""

from __future__ import annotations

from datetime import date, datetime
import os
from pathlib import Path
import re
import threading
from typing import Optional


DATA_DIR = Path(".")


def configure_checklist(data_dir: Path) -> None:
    """Bind the portable data directory without coupling this module to FastAPI."""
    global DATA_DIR
    DATA_DIR = data_dir


# Maps month-folder names ("Jan 2026", "April 2026", ...) to a number so months
# sort newest-first regardless of full/abbreviated spelling.
MONTH_ORDER = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}

# Synthetic example tree for local demos. In real use the dashboard derives this
# from the month's checklist.xlsx instead of inferring coverage only from files
# that are currently present.
SYSTEM_TREE = [
    {
        "system": "RADIO",
        "locations": [
            {"name": "HQ", "expected": 4, "antivirus": "ClamWin"},
            {"name": "HQ East", "expected": 2, "antivirus": "ClamWin"},
            {"name": "Site A", "expected": 2, "antivirus": "ClamWin"},
            {"name": "Depot", "expected": 2, "antivirus": "ClamWin"},
        ],
    },
    {
        "system": "CCTV",
        "locations": [
            {"name": "HQ", "expected": 6, "antivirus": "McAfee"},
            {"name": "Operations Center", "expected": 6, "antivirus": "McAfee"},
            {"name": "Site A", "expected": 4, "antivirus": "McAfee"},
            {"name": "Site B", "expected": 4, "antivirus": "McAfee"},
            {"name": "Site C", "expected": 3, "antivirus": "McAfee"},
        ],
    },
    {
        "system": "EPM",
        "locations": [
            {"name": "Depot", "expected": 6, "antivirus": "McAfee"},
            {"name": "Site A", "expected": 2, "antivirus": "McAfee"},
            {"name": "Site B", "expected": 2, "antivirus": "McAfee"},
        ],
    },
    {
        "system": "INFO",
        "locations": [
            {"name": "HQ", "expected": 5, "antivirus": "McAfee"},
            {"name": "Site A", "expected": 3, "antivirus": "McAfee"},
            {"name": "Site C", "expected": 2, "antivirus": "McAfee"},
        ],
    },
    {
        "system": "NET",
        "locations": [
            {"name": "HQ East", "expected": 3, "antivirus": "Trellix"},
            {"name": "Data Center", "expected": 1, "antivirus": "Trellix"},
        ],
    },
    {
        "system": "VOICE",
        "locations": [
            {"name": "Site U", "expected": 3, "antivirus": "Trellix"},
            {"name": "Ops Center", "expected": 2, "antivirus": "Trellix"},
        ],
    },
    {
        "system": "GW",
        "locations": [
            {"name": "Ops Center", "expected": 6, "antivirus": "Unknown"},
            {"name": "HQ", "expected": 6, "antivirus": "Unknown"},
            {"name": "Site A", "expected": 2, "antivirus": "Unknown"},
        ],
    },
]

CHECKLIST_SHEET_SYSTEMS = {
    "RADIO": "RADIO",
    "SUPERVISORY": "GW",
}
CHECKLIST_SYSTEM_ORDER = []
CHECKLIST_ROWS = []
CHECKLIST_CACHE = {}
checklist_cache_lock = threading.Lock()
CHECKLIST_EXTENSIONS = {'.xlsx', '.xlsm'}

DEFAULT_ANTIVIRUS_BY_SYSTEM = {
    "RADIO": "ClamWin",
    "CCTV": "McAfee",
    "EPM": "McAfee",
    "INFO": "McAfee",
    "NET": "Trellix",
    "PBX": "Trellix",
    "AUDIO": "Trellix",
    "VOICE": "Trellix",
    "FAC": "Trellix",
    "NVR": "Trellix",
    "FW": "Trellix",
}

COMMON_MATCH_TOKENS = {
    "pc", "client", "server", "system", "radio", "trellix", "mcafee", "clamwin",
    "nms", "user", "scan", "done", "yes", "no", "the", "and", "for", "rom",
}

def checklist_month_tokens(month: Optional[str]):
    if not month:
        return set()

    aliases = {
        "january": {"january", "jan"},
        "february": {"february", "feb"},
        "march": {"march", "mar"},
        "april": {"april", "apr"},
        "may": {"may"},
        "june": {"june", "jun"},
        "july": {"july", "jul"},
        "august": {"august", "aug"},
        "september": {"september", "sep", "sept"},
        "october": {"october", "oct"},
        "november": {"november", "nov"},
        "december": {"december", "dec"},
    }

    tokens = set()
    for token in re.findall(r'[a-zA-Z]+|\d{4}', str(month).lower()):
        tokens.add(token)
        for month_name, month_aliases in aliases.items():
            if token == month_name or token in month_aliases:
                tokens.update(month_aliases)
    return tokens

def is_checklist_workbook(path: Path):
    return (
        path.is_file()
        and path.suffix.lower() in CHECKLIST_EXTENSIONS
        and not path.name.startswith("~$")
        and not path.name.startswith(".")
    )

def checklist_candidate_score(path: Path, month: Optional[str]):
    name = path.name.lower()
    score = 0
    if "checklist" in name:
        score += 100
    if "update" in name or "updated" in name:
        score += 30
    if "with" in name:
        score += 5
    if "antivirus" in name or "av" in name or "status" in name:
        score += 15
    if "log" in name:
        score += 5
    for token in checklist_month_tokens(month):
        if token and token in name:
            score += 8
    try:
        file_stat = path.stat()
        return (score, file_stat.st_mtime, file_stat.st_size, path.name.lower())
    except OSError:
        return (score, 0, 0, path.name.lower())

def checklist_path(month: Optional[str] = None):
    month_dir = DATA_DIR / month if month else None
    if month_dir and month_dir.exists():
        # ``os.scandir`` uses the directory entry metadata where available. It
        # avoids an extra ``stat`` call for every log in a large month folder.
        candidates = []
        try:
            with os.scandir(month_dir) as entries:
                for entry in entries:
                    name = entry.name
                    if (
                        name.startswith(("~$", "."))
                        or Path(name).suffix.lower() not in CHECKLIST_EXTENSIONS
                    ):
                        continue
                    try:
                        if entry.is_file():
                            candidates.append(Path(entry.path))
                    except OSError:
                        continue
        except OSError:
            candidates = []
        checklist_candidates = [path for path in candidates if "checklist" in path.name.lower()]
        if checklist_candidates:
            return max(checklist_candidates, key=lambda path: checklist_candidate_score(path, month))
        if candidates:
            return max(candidates, key=lambda path: checklist_candidate_score(path, month))

    if month:
        # Never validate one month against another month's workbook.  A missing
        # month checklist is an explicit reconciliation error, not a fallback.
        return (DATA_DIR / month / f"{month} checklist.xlsx")

    default_path = DATA_DIR / "default checklist.xlsx"
    if default_path.exists():
        return default_path

    preferred = DATA_DIR / "Jun 2026" / "June 2026 checklist.xlsx"
    if preferred.exists():
        return preferred
    matches = sorted(
        (path for path in DATA_DIR.glob("*/June 2026 checklist.xlsx") if is_checklist_workbook(path)),
        key=lambda path: str(path).lower(),
    )
    return matches[0] if matches else preferred

def cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def normalize_checklist_header(value):
    return re.sub(r'[^a-z0-9]', '', cell_to_text(value).lower())

def normalize_lookup(value):
    return re.sub(r'[^a-z0-9]', '', str(value).lower())

def checklist_tokens(*values):
    tokens = set()
    for value in values:
        raw = str(value or "").lower()
        compact = normalize_lookup(raw)
        if len(compact) >= 3:
            tokens.add(compact)
        for token in re.findall(r'[a-z0-9]+', raw):
            if len(token) >= 2 and token not in COMMON_MATCH_TOKENS:
                tokens.add(token)
    return tokens

def infer_checklist_system(sheet_name: str, row_system: str):
    sheet_key = sheet_name.strip().upper()
    if sheet_key.startswith("RADIO"):
        mapped = "RADIO"
    elif sheet_key.startswith("SUPERVISORY"):
        mapped = "GW"
    else:
        mapped = CHECKLIST_SHEET_SYSTEMS.get(sheet_key, sheet_key)

    row_key = row_system.strip().upper()
    if mapped == "NET" and row_key.startswith("FW"):
        return "FW"
    return mapped

def normalize_antivirus(value: str, system_name: str):
    text = cell_to_text(value)
    if not text:
        return DEFAULT_ANTIVIRUS_BY_SYSTEM.get(system_name, "Unknown")
    lower = text.lower()
    if "clam" in lower:
        return "ClamWin"
    if "mcafee" in lower or "epo" in lower:
        return "McAfee"
    if "trel" in lower or "trellix" in lower:
        return "Trellix"
    return text

def find_header_row(ws):
    """Find headers with one sequential read of a read-only worksheet."""
    for row_index, values in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True),
        start=1,
    ):
        headers = [normalize_checklist_header(value) for value in values]
        if "location" in headers and "system" in headers:
            return row_index, headers
    return None, []

def first_header_index(headers, *names):
    for name in names:
        normalized = normalize_checklist_header(name)
        if normalized in headers:
            return headers.index(normalized) + 1
    return None

def empty_checklist_data(path: Optional[Path] = None, error: str = ""):
    return {
        "rows": [],
        "system_order": [],
        "path": str(path) if path else "",
        "valid": False,
        "error": error or "The monthly checklist is unavailable.",
    }

def checklist_row_identity(row):
    return (
        normalize_lookup(row.get("system", "")),
        normalize_lookup(row.get("location", "")),
        normalize_lookup(row.get("item", "")),
        normalize_lookup(row.get("pc_name", "")),
    )

def checklist_cell(values, column):
    """Safely read a 1-based checklist column from a streamed row tuple."""
    if not column or column > len(values):
        return ""
    return values[column - 1]

def load_checklist_rows(month: Optional[str] = None):
    path = checklist_path(month)
    if not path.exists() or not is_checklist_workbook(path):
        return empty_checklist_data(path, "No valid checklist workbook was found for this month.")
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        return empty_checklist_data(
            path,
            f"The checklist workbook could not be parsed safely ({type(exc).__name__}).",
        )

    rows = []
    system_order = []
    seen_sheet_signatures = set()
    for ws in workbook.worksheets:
        if ws.title.strip().lower() == "global checklist":
            continue
        header_row, headers = find_header_row(ws)
        if not header_row:
            continue

        location_col = first_header_index(headers, "Location")
        system_col = first_header_index(headers, "System")
        pc_col = first_header_index(headers, "PC Name", "Host Name", "Hostname", "System Name")
        dat_col = first_header_index(headers, "DAT version", "DAT Version")
        date_col = first_header_index(headers, "Date Scan", "Date Scanned")
        scan_done_col = first_header_index(headers, "Scan Done")
        virus_found_col = first_header_index(headers, "Virus Found")
        screenshot_col = first_header_index(headers, "Screenshot Upload")
        done_by_col = first_header_index(headers, "Done By")
        endorsed_by_col = first_header_index(headers, "Endorsed By", "Endorsed By Eric")
        av_col = first_header_index(headers, "AV Software")
        time_col = first_header_index(headers, "Time Scanning")
        remarks_cols = [index + 1 for index, header in enumerate(headers) if "remarks" in header]
        sheet_rows = []

        # Read-only worksheets are forward-only. Calling ``ws.cell`` for every
        # field restarts their XML stream repeatedly, so one large checklist can
        # take minutes. Streaming each row once preserves every source value
        # while making checklist loading linear in the number of rows.
        for row_index, values in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            location = cell_to_text(checklist_cell(values, location_col))
            row_system = cell_to_text(checklist_cell(values, system_col))
            if not location and not row_system:
                continue

            system_name = infer_checklist_system(ws.title, row_system)
            if system_name not in system_order:
                system_order.append(system_name)

            pc_name = cell_to_text(checklist_cell(values, pc_col))
            dat_version = cell_to_text(checklist_cell(values, dat_col))
            date_scan = cell_to_text(checklist_cell(values, date_col))
            remarks = "; ".join(
                text for text in (cell_to_text(checklist_cell(values, col)) for col in remarks_cols)
                if text
            )
            av_software = normalize_antivirus(checklist_cell(values, av_col), system_name)
            search_values = [system_name, location, row_system, pc_name, ws.title]
            row = {
                "id": f"{ws.title}:{row_index}",
                "sheet": ws.title,
                "row": row_index,
                "system": system_name,
                "location": location,
                "item": row_system,
                "pc_name": pc_name,
                "dat_version": dat_version,
                "date_scan": date_scan,
                "scan_done": cell_to_text(checklist_cell(values, scan_done_col)),
                "virus_found": cell_to_text(checklist_cell(values, virus_found_col)),
                "screenshot_upload": cell_to_text(checklist_cell(values, screenshot_col)),
                "done_by": cell_to_text(checklist_cell(values, done_by_col)),
                "endorsed_by": cell_to_text(checklist_cell(values, endorsed_by_col)),
                "av_software": av_software,
                "time_scanning": cell_to_text(checklist_cell(values, time_col)),
                "remarks": remarks,
                "lookup": normalize_lookup(" ".join(search_values)),
                "tokens": sorted(checklist_tokens(*search_values)),
            }
            sheet_rows.append(row)

        if not sheet_rows:
            continue
        sheet_signature = tuple(checklist_row_identity(row) for row in sheet_rows)
        if sheet_signature in seen_sheet_signatures:
            continue
        seen_sheet_signatures.add(sheet_signature)
        rows.extend(sheet_rows)

    try:
        workbook.close()
    except Exception:
        pass

    valid = bool(rows)
    return {
        "rows": rows,
        "system_order": system_order,
        "path": str(path),
        "valid": valid,
        "error": "" if valid else "The checklist contains no supported evidence rows.",
    }

def build_system_tree_from_checklist(rows, system_order=None):
    if not rows:
        return []
    grouped = {}
    location_order = {}
    antivirus_counts = {}
    for row in rows:
        system_name = row["system"]
        location = row["location"] or "Unspecified"
        grouped.setdefault(system_name, {})
        location_order.setdefault(system_name, [])
        antivirus_counts.setdefault(system_name, {})
        grouped[system_name][location] = grouped[system_name].get(location, 0) + 1
        if location not in location_order[system_name]:
            location_order[system_name].append(location)
        antivirus = row.get("av_software") or DEFAULT_ANTIVIRUS_BY_SYSTEM.get(system_name, "Unknown")
        antivirus_counts[system_name].setdefault(location, {})
        antivirus_counts[system_name][location][antivirus] = antivirus_counts[system_name][location].get(antivirus, 0) + 1

    ordered_systems = [system for system in (system_order or CHECKLIST_SYSTEM_ORDER) if system in grouped]
    ordered_systems.extend(system for system in grouped if system not in ordered_systems)
    tree = []
    for system_name in ordered_systems:
        locations = []
        for location in location_order.get(system_name, sorted(grouped[system_name])):
            counts = antivirus_counts.get(system_name, {}).get(location, {})
            antivirus = max(counts, key=counts.get) if counts else DEFAULT_ANTIVIRUS_BY_SYSTEM.get(system_name, "Unknown")
            locations.append({
                "name": location,
                "expected": grouped[system_name][location],
                "antivirus": antivirus,
            })
        tree.append({"system": system_name, "locations": locations})
    return tree

def get_month_checklist(month: Optional[str] = None):
    path = checklist_path(month)
    try:
        file_stat = path.stat()
        cache_signature = (str(path.resolve()), file_stat.st_mtime_ns, file_stat.st_size)
    except OSError:
        cache_signature = (str(path), 0, 0)

    cache_key = month or "__default__"
    with checklist_cache_lock:
        cached = CHECKLIST_CACHE.get(cache_key)
        if cached and cached.get("cache_signature") == cache_signature:
            return cached

    loaded = load_checklist_rows(month)
    rows = loaded.get("rows", [])
    system_order = loaded.get("system_order", [])
    data = {
        "rows": rows,
        "system_order": system_order,
        "path": loaded.get("path", str(path)),
        "valid": bool(loaded.get("valid", False)),
        "error": str(loaded.get("error", "")),
        "system_tree": build_system_tree_from_checklist(rows, system_order),
        "cache_signature": cache_signature,
    }

    with checklist_cache_lock:
        CHECKLIST_CACHE[cache_key] = data
    return data

def extract_signature_number(value: str):
    if not value:
        return None
    match = re.search(r'\d+(?:\.\d+)?', str(value).replace(',', '.'))
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None

def parse_date_candidates(value: str):
    text = str(value or "")
    candidates = set()
    iso_match = re.search(r'\b(\d{4})-(\d{1,2})-(\d{1,2})\b', text)
    if iso_match:
        year, month, day = map(int, iso_match.groups())
        try:
            candidates.add(date(year, month, day))
        except ValueError:
            pass
    slash_match = re.search(r'\b(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\b', text)
    if slash_match:
        first, second, year = map(int, slash_match.groups())
        if year < 100:
            year += 2000
        for month, day in ((first, second), (second, first)):
            try:
                candidates.add(date(year, month, day))
            except ValueError:
                pass
    text_match = re.search(r'\b(?:Mon|Tue|Wed|Thu|Fri|Sat|Sun)?\s*([A-Za-z]{3,9})\s+(\d{1,2})\s+\d{1,2}:\d{2}:\d{2}\s+(\d{4})\b', text, re.IGNORECASE)
    if text_match:
        month_text, day, year = text_match.groups()
        for fmt in ("%B", "%b"):
            try:
                month = datetime.strptime(month_text[:3], "%b").month if fmt == "%b" else datetime.strptime(month_text, "%B").month
                candidates.add(date(int(year), month, int(day)))
                break
            except ValueError:
                continue
    return candidates

def dates_match(expected: str, actual: str):
    expected_dates = parse_date_candidates(expected)
    actual_dates = parse_date_candidates(actual)
    return bool(expected_dates and actual_dates and expected_dates.intersection(actual_dates))

def _checklist_location_hint(system_name: str, subsystem_name: str, filename: str):
    """Translate known legacy radio folders into checklist locations."""
    if system_name != "RADIO":
        return subsystem_name
    compact_subsystem = normalize_lookup(subsystem_name)
    compact_filename = normalize_lookup(filename)
    if compact_subsystem == "siteradio":
        site = re.match(r"\s*([A-Za-z]{3})\b", filename)
        return site.group(1).upper() if site else subsystem_name
    if compact_subsystem == "depot":
        if re.search(r"\bas\s*[-_]?\s*[12]\b", filename, re.IGNORECASE):
            return "Operations Center"
        if any(token in compact_filename for token in ("crs", "gw", "nms")):
            return "HQ East"
    return subsystem_name


def _equipment_key(value: str):
    """Normalise checklist equipment labels without losing unit numbers."""
    tokens = re.findall(r"[a-z]+|\d+", str(value or "").lower())
    while tokens and tokens[0] in {
        "radio", "cctv", "epm", "info", "net", "pbx", "audio",
        "voice", "fac", "nvr", "fw", "gw", "supervisory",
    }:
        tokens.pop(0)
    return "".join(tokens)


def _same_checklist_target(left, right):
    return (
        normalize_lookup(left.get("location", "")),
        normalize_lookup(left.get("item", "")),
        normalize_lookup(left.get("pc_name", "")),
    ) == (
        normalize_lookup(right.get("location", "")),
        normalize_lookup(right.get("item", "")),
        normalize_lookup(right.get("pc_name", "")),
    )


def checklist_match_score(log_entry, subsystem_name: str, row):
    """Score one evidence/checklist pair without assigning the row twice."""
    system_name = log_entry.get("system", "")
    if row.get("system") != system_name:
        return 0

    raw_filename = Path(log_entry.get("path", log_entry.get("filename", ""))).stem
    display_filename = log_entry.get("filename", "")
    location_hint = _checklist_location_hint(system_name, subsystem_name, raw_filename)
    log_lookup = normalize_lookup(" ".join([system_name, location_hint, raw_filename, display_filename]))
    log_tokens = checklist_tokens(system_name, subsystem_name, raw_filename, display_filename)
    normalized_subsystem = normalize_lookup(location_hint)
    normalized_stem = normalize_lookup(raw_filename)
    score = 0

    location_lookup = normalize_lookup(row.get("location", ""))
    if location_lookup and location_lookup == normalized_subsystem:
        score += 16
    elif location_lookup and location_lookup in log_lookup:
        score += 10

    for value in (row.get("item", ""), row.get("pc_name", "")):
        compact = normalize_lookup(value)
        equipment = _equipment_key(value)
        if equipment and equipment == normalized_stem:
            score += 48
        elif equipment and equipment in normalized_stem:
            score += 36
        else:
            unit_match = re.fullmatch(r"([a-z]+)(\d+)", equipment)
            if unit_match and unit_match.group(1) in normalized_stem and normalized_stem.endswith(unit_match.group(2)):
                score += 24
        if compact and compact == normalized_stem:
            score += 60
        elif compact and compact in log_lookup:
            score += 24
        elif compact and log_lookup and log_lookup in compact:
            score += 8
        if compact.startswith(normalize_lookup(system_name)):
            trimmed = compact[len(normalize_lookup(system_name)):]
            if trimmed and trimmed in log_lookup:
                score += 18

    common_tokens = log_tokens.intersection(set(row.get("tokens", [])))
    score += min(len(common_tokens), 6) * 3
    return score


def match_checklist_row(log_entry, subsystem_name: str, checklist_rows=None):
    """Find the best checklist row for one evidence entry using the shared scorer."""
    rows = CHECKLIST_ROWS if checklist_rows is None else checklist_rows
    system_name = log_entry.get("system", "")
    candidates = [row for row in rows if row.get("system") == system_name]
    if not candidates:
        return None

    scored = [(checklist_match_score(log_entry, subsystem_name, row), row) for row in candidates]
    scored.sort(key=lambda pair: (-pair[0], str(pair[1].get("id", ""))))
    best_score, best_row = scored[0]
    return best_row if best_score >= 8 else None


def match_checklist_rows(log_entries, checklist_rows=None):
    """Assign checklist rows once, surfacing ties and contested evidence."""
    rows = CHECKLIST_ROWS if checklist_rows is None else checklist_rows
    assignments = [None] * len(log_entries)
    metadata = [
        {"score": 0, "ambiguous": False, "reason": "No checklist row matched."}
        for _entry in log_entries
    ]
    pairs = []

    for log_index, log_entry in enumerate(log_entries):
        subsystem_name = log_entry.get("subsystem", "Base System")
        scored = []
        for row_index, row in enumerate(rows):
            score = checklist_match_score(log_entry, subsystem_name, row)
            if score >= 8:
                scored.append((score, row_index))
        scored.sort(key=lambda value: (-value[0], str(rows[value[1]].get("id", value[1]))))
        if not scored:
            continue
        tied_best = [row_index for score, row_index in scored if score == scored[0][0]]
        if (
            len(tied_best) > 1
            and not all(_same_checklist_target(rows[tied_best[0]], rows[row_index]) for row_index in tied_best[1:])
        ):
            metadata[log_index] = {
                "score": scored[0][0],
                "ambiguous": True,
                "reason": "Multiple checklist rows matched with the same confidence.",
            }
            continue
        for score, row_index in scored:
            pairs.append((
                -score,
                str(log_entry.get("path", "")).lower(),
                str(rows[row_index].get("id", row_index)),
                log_index,
                row_index,
            ))

    assigned_rows = set()
    for negative_score, _path, _row_id, log_index, row_index in sorted(pairs):
        if assignments[log_index] is not None or row_index in assigned_rows:
            continue
        assignments[log_index] = rows[row_index]
        assigned_rows.add(row_index)
        metadata[log_index] = {
            "score": -negative_score,
            "ambiguous": False,
            "reason": "",
        }

    for log_index, assignment in enumerate(assignments):
        if assignment is None and any(pair[3] == log_index for pair in pairs):
            best_score = max(-pair[0] for pair in pairs if pair[3] == log_index)
            metadata[log_index] = {
                "score": best_score,
                "ambiguous": True,
                "reason": "The best checklist row was already assigned to other evidence.",
            }
    return assignments, metadata

def build_checklist_validation(log_entry, checklist_row, status: str):
    """Compare checklist values without changing the evidence verdict.

    Checklist differences are reconciliation metadata, not scan failures.
    """
    fields = {
        "checklist_remarks": "",
        "checklist_dat_version": "",
        "checklist_date_scan": "",
        "checklist_virus_found": "",
        "checklist_av_software": "",
        "checklist_match": "No checklist row",
        "checklist_mismatches": [],
    }
    if not checklist_row:
        return status, fields

    fields.update({
        "checklist_remarks": checklist_row.get("remarks", ""),
        "checklist_dat_version": checklist_row.get("dat_version", ""),
        "checklist_date_scan": checklist_row.get("date_scan", ""),
        "checklist_virus_found": checklist_row.get("virus_found", ""),
        "checklist_av_software": checklist_row.get("av_software", ""),
        "checklist_match": f"{checklist_row.get('sheet')} row {checklist_row.get('row')}",
    })

    mismatches = []
    checklist_dat = checklist_row.get("dat_version", "")
    log_dat = log_entry.get("dat_version", "")
    if status != "Not Scanned" and checklist_dat:
        checklist_number = extract_signature_number(checklist_dat)
        log_number = extract_signature_number(log_dat)
        if log_dat == "Unknown" or log_number is None:
            mismatches.append("DAT version not detected in scanned log")
        elif checklist_number is not None:
            if not (log_entry.get("system") == "RADIO" and checklist_number < 1000 <= log_number):
                if abs(checklist_number - log_number) >= 0.01:
                    mismatches.append(f"DAT mismatch: checklist {checklist_dat}, scanned {log_dat}")

    checklist_date = checklist_row.get("date_scan", "")
    log_scan = log_entry.get("last_scan", "")
    if status != "Not Scanned" and checklist_date:
        if log_scan == "Unknown":
            mismatches.append("Scan date not detected in scanned log")
        elif not dates_match(checklist_date, log_scan):
            mismatches.append(f"Scan date mismatch: checklist {checklist_date}, scanned {log_scan}")

    virus_found = checklist_row.get("virus_found", "").strip().lower()
    if virus_found in {"no", "none", "0"} and status == "Threats Found":
        mismatches.append("Checklist says no virus found, scanned log reports threat")
    elif virus_found in {"yes", "y", "true"} and status == "Clean":
        mismatches.append("Checklist says virus found, scanned log is clean")

    fields["checklist_mismatches"] = mismatches
    return status, fields
